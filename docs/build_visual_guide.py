from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
UPLOADS = Path('/workspace/scratch/e64c1d57c9ca/upload')
ASSETS = ROOT / 'visual-assets'
OUT = ROOT / 'TIME-Escola-de-Negocios-Guia-Visual-para-Socios.docx'
ASSETS.mkdir(exist_ok=True)

NAVY = '050829'; VIOLET = '3700F0'; MAGENTA = 'FF00E0'; INK = '18172D'; MUTED = '625F76'; PALE = 'F5F4FA'; BORDER = 'DEDBEA'; TEAL = '087B68'

def crop(name, source, top, bottom):
    target = ASSETS / name
    im = Image.open(UPLOADS / source)
    w, h = im.size
    im.crop((0, max(0, top), w, min(h, bottom))).save(target, quality=92)
    return target

IMAGES = {
    'landing': crop('landing.png', 'D19DFD6F-550B-49D8-8EB1-B65148BFD74B.png', 480, 1420),
    'onboarding': crop('onboarding.png', '0E604969-F408-4AA2-868C-817C9316473F.png', 335, 1660),
    'courses': crop('courses.png', 'D19DFD6F-550B-49D8-8EB1-B65148BFD74B.png', 335, 1640),
    'connections': crop('connections.png', '197723ab-8e77-4c55-8f42-77ee9c282271.png', 0, 1250),
    'community': crop('community.png', '0FAE329A-902D-4369-9FEF-694196A15515.png', 335, 1645),
}

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def set_cell_margin(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); tcMar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_border(cell, color=BORDER):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        node = OxmlElement(f'w:{edge}'); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'8'); node.set(qn('w:color'),color); borders.append(node)

def run_style(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = 'Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic

def para(doc_or_cell, text='', size=11, color=INK, bold=False, italic=False, align=None, before=0, after=6, line=1.14):
    p = doc_or_cell.add_paragraph() if hasattr(doc_or_cell, 'add_paragraph') else doc_or_cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = line
    if align is not None: p.alignment = align
    r = p.add_run(text); run_style(r, size, color, bold, italic); return p

def reset_first_paragraph(cell):
    p = cell.paragraphs[0]; p.clear(); p.paragraph_format.space_after = Pt(0); return p

def add_kicker(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7); r=p.add_run(text.upper()); run_style(r, 9, MAGENTA, True); r.font.letter_spacing = Pt(1.2) if hasattr(r.font,'letter_spacing') else None
    return p

def add_title(doc, text, subtitle=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); r=p.add_run(text); run_style(r, 28, NAVY, True)
    if subtitle: para(doc, subtitle, 12, MUTED, after=18, line=1.22)

def page_break(doc): doc.add_page_break()

def add_image(doc, path, width=2.55, align=WD_ALIGN_PARAGRAPH.CENTER):
    p=doc.add_paragraph(); p.alignment=align; p.paragraph_format.space_after=Pt(4); p.add_run().add_picture(str(path), width=Inches(width)); return p

def caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(13); r=p.add_run(text); run_style(r, 9.3, MUTED, italic=True)

def add_callout(doc, title, body, accent=VIOLET):
    table=doc.add_table(rows=1, cols=1); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False; cell=table.cell(0,0); shade(cell, PALE); set_border(cell, 'DDD9EE'); set_cell_margin(cell, 140, 180, 140, 180)
    p=reset_first_paragraph(cell); r=p.add_run(title); run_style(r, 11, accent, True); p.paragraph_format.space_after=Pt(4)
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(body); run_style(r,10.5,INK); return table

def add_steps(doc, steps):
    table=doc.add_table(rows=len(steps), cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths=[Inches(.55), Inches(5.95)]
    for row, (number, title, body) in zip(table.rows, steps):
        for cell, width in zip(row.cells, widths): cell.width=width; set_cell_margin(cell,100,130,100,130); set_border(cell, 'E4E0EF')
        shade(row.cells[0], 'F0ECFF'); row.cells[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=reset_first_paragraph(row.cells[0]); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(number); run_style(r, 13, VIOLET, True)
        p=reset_first_paragraph(row.cells[1]); r=p.add_run(title+'\n'); run_style(r, 11, NAVY, True); r=p.add_run(body); run_style(r,10.2,MUTED)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def add_feature_page(doc, kicker, title, intro, image, caption_text, steps, note=None):
    page_break(doc); add_kicker(doc, kicker); add_title(doc, title, intro); add_image(doc, image); caption(doc, caption_text); add_steps(doc, steps)
    if note: add_callout(doc, note[0], note[1], note[2] if len(note)>2 else VIOLET)

doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11); section.top_margin=Inches(.72); section.bottom_margin=Inches(.72); section.left_margin=Inches(.88); section.right_margin=Inches(.88)
section.header_distance=Inches(.3); section.footer_distance=Inches(.35)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal'].font.size=Pt(10.5)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(54); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('TIME'); run_style(r,42,NAVY,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22); r=p.add_run('ESCOLA DE NEGÓCIOS'); run_style(r,11,MUTED,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8); r=p.add_run('Guia visual da plataforma'); run_style(r,30,NAVY,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(26); r=p.add_run('Para apresentação aos sócios'); run_style(r,14,MUTED)
add_image(doc, IMAGES['landing'], 5.9)
caption(doc, 'Visão institucional da TIME: conhecimento, conexões, eventos e negócios em uma experiência única.')
add_callout(doc, 'O QUE A PLATAFORMA RESOLVE', 'A TIME organiza aprendizado, empresários, oportunidades e eventos em um único ambiente. O foco do MVP é transformar dados reais dos membros em conexões úteis.', MAGENTA)

page_break(doc); add_kicker(doc,'VISÃO GERAL'); add_title(doc,'Uma jornada simples para o membro', 'A pessoa entra com um cadastro curto e completa o perfil aos poucos, conforme começa a usar a plataforma.')
add_steps(doc, [('1','Criar conta','Nome, e-mail e senha. O primeiro acesso não começa com um formulário longo.'),('2','Conhecer a plataforma','O membro vê a área interna, cursos em preparação, conexões e avisos.'),('3','Completar o perfil','A nota de completude mostra exatamente o que falta para aproveitar as áreas liberadas.'),('4','Aprender e se conectar','Cursos, comunidade, diretório e eventos passam a trabalhar juntos.')])
add_callout(doc,'IMPORTANTE','A plataforma não inventa membros, cursos, eventos ou resultados. Quando a TIME ainda não publicou um conteúdo, ela mostra “Em breve” ou um estado vazio claro.',TEAL)

add_feature_page(doc, 'ENTRADA E PERFIL', 'Cadastro progressivo, sem cansar no primeiro acesso.', 'O perfil é construído em etapas curtas. A nota de 0 a 100 deixa claro o que melhora a presença do empresário na comunidade.', IMAGES['onboarding'], 'Tela de completude: percentual, progresso e itens que o membro ainda precisa informar.', [('1','Entrar e explorar','O membro pode conhecer a plataforma logo depois do login.'),('2','Ver a nota do perfil','A nota mostra quanto já foi preenchido, sem punir ou confundir.'),('3','Completar o que importa','Foto, empresa, profissão, segmento, cidade, WhatsApp, o que oferece e o que procura fortalecem as conexões.')], ('LIBERAÇÃO PROGRESSIVA','Cursos completos, comunidade e networking dependem do perfil com informações reais e revisadas.',MAGENTA))

add_feature_page(doc, 'CONHECIMENTO', 'Cursos em uma experiência de biblioteca.', 'A área de cursos foi desenhada como uma biblioteca visual. Conteúdos ainda não confirmados aparecem como “Em breve”, sem inventar professores ou aulas.', IMAGES['courses'], 'Tela de cursos: busca, tema, progresso real e estado honesto enquanto a biblioteca oficial é preparada.', [('1','Encontrar conteúdo','Busca e categorias ajudam o membro a achar a próxima etapa da jornada.'),('2','Continuar assistindo','Quando houver aulas publicadas, o progresso é salvo e a pessoa retoma de onde parou.'),('3','Aprender dentro da TIME','O vídeo é exibido dentro da plataforma, com título, módulo, descrição e botão de conclusão.')], ('VÍDEOS FORA DA TIME','A equipe cola o link de um vídeo não listado do YouTube, Vimeo ou Panda. A TIME não hospeda o arquivo de vídeo; apenas incorpora o player.',VIOLET))

add_feature_page(doc, 'CONEXÕES', 'O diretório transforma necessidade em conversa.', 'O membro descreve o profissional, empresa ou serviço que procura. A busca considera somente perfis reais cadastrados na comunidade.', IMAGES['connections'], 'Tela do diretório: busca por nome, empresa, segmento, especialidade ou serviço.', [('1','Descrever a necessidade','Ex.: “preciso de tráfego pago para uma clínica”.'),('2','Encontrar perfis reais','A plataforma cruza a necessidade com o que os membros informaram em seus perfis.'),('3','Abrir o perfil e chamar','O membro entende por que a indicação faz sentido e inicia a conversa pelo WhatsApp.')], ('REGRA DE CONFIANÇA','Nenhuma recomendação deve criar pessoas, empresas, cidades ou especialidades que não existam no banco da TIME.',TEAL))

add_feature_page(doc, 'COMUNIDADE', 'Avisos, aprendizados e oportunidades no mesmo lugar.', 'A Comunidade é o espaço para compartilhar uma conquista, uma dúvida ou uma oportunidade de negócio, sem depender de grupos externos.', IMAGES['community'], 'Tela da Comunidade: campo de publicação, posts, navegação inferior e identificação do membro.', [('1','Publicar um contexto','Um membro compartilha uma conquista, dúvida ou oportunidade.'),('2','Acompanhar os avisos','A TIME pode utilizar a área para comunicados relevantes aos membros.'),('3','Manter a conversa contextual','A comunidade reforça as conexões que começam por interesses e necessidades reais.')], ('PRIVACIDADE','Informações de contato não devem aparecer publicamente sem consentimento. O WhatsApp é acessado pelo botão de contato do perfil.',MAGENTA))

page_break(doc); add_kicker(doc,'EVENTOS E EXPERIÊNCIAS'); add_title(doc,'Da agenda ao encontro presencial.', 'Os eventos aparecem na plataforma e também podem ser encontrados pela home. Cada evento tem informações reais de data, local e inscrição.')
add_steps(doc,[('1','Ver próximos eventos','A agenda mostra apenas eventos publicados pela TIME.'),('2','Abrir os detalhes','O membro acessa descrição, local, data, horário e orientações.'),('3','Confirmar presença','A inscrição é registrada para a equipe acompanhar os participantes.'),('4','Adicionar à agenda','O convite pode direcionar o membro para o calendário do celular.')])
add_callout(doc,'O QUE NÃO SERÁ PROMETIDO','Enquanto não houver evento publicado pela TIME, a área deve permanecer com um estado vazio claro. Não usar imagens, datas ou participantes fictícios.',TEAL)

page_break(doc); add_kicker(doc,'ADMINISTRAÇÃO'); add_title(doc,'O painel que mantém a plataforma viva.', 'Somente uma conta com papel administrativo real pode gerenciar o conteúdo. A autorização é validada no servidor e no banco; não depende apenas do e-mail digitado.')
add_steps(doc,[('1','Membros e empresas','Pesquisar, revisar e acompanhar a completude dos perfis.'),('2','Cursos e aulas','Criar cursos, adicionar capa, módulos, título da aula, descrição e link externo de vídeo.'),('3','Eventos e inscritos','Publicar uma agenda, receber inscrições e acompanhar a lista de participantes.'),('4','Comunidade e avisos','Administrar comunicados e manter o ambiente útil para os membros.'),('5','Configurações','Organizar dados institucionais, canais oficiais e integrações autorizadas.')])
add_callout(doc,'DEMONSTRAÇÃO','O acesso de demonstração deve ser usado apenas pelos sócios autorizados. Ele precisa ser criado como usuário real com papel admin, nunca como um atalho baseado em e-mail.',MAGENTA)

page_break(doc); add_kicker(doc,'COMO INSTALAR'); add_title(doc,'TIME no celular: PWA instalável.', 'A TIME funciona como uma PWA: uma experiência instalável no celular, sem afirmar que já existe aplicativo nativo na App Store ou Play Store.')
add_steps(doc,[('iOS','No iPhone','Abra a TIME no Safari, toque em Compartilhar e escolha “Adicionar à Tela de Início”.'),('Android','No Android','Abra no Chrome e use “Instalar app” quando o navegador oferecer a opção.'),('Uso','Depois de instalar','O ícone abre a TIME em uma experiência de aplicativo, com navegação mobile e área segura para o conteúdo.')])
add_callout(doc,'LINK DE APRESENTAÇÃO','https://time-escola-negocios.ismaelcharello.chatgpt.site',VIOLET)

page_break(doc); add_kicker(doc,'ROTEIRO PARA OS SÓCIOS'); add_title(doc,'Demonstração de 8 minutos', 'Uma apresentação enxuta, sem prometer recursos que ainda estão em evolução.')
add_steps(doc,[('1','Começo — 1 min','Abrir a landing e explicar: conhecimento que gera conexões; conexões que geram negócios.'),('2','Entrada — 1 min','Mostrar login/criar conta e explicar o cadastro progressivo.'),('3','Perfil — 1 min','Mostrar a nota de completude e o que um empresário precisa informar.'),('4','Cursos — 1 min','Mostrar a biblioteca, “Em breve” honesto e como os vídeos externos serão exibidos.'),('5','Conexões — 2 min','Fazer uma busca e abrir um perfil, explicando a recomendação e o contato por WhatsApp.'),('6','Comunidade e eventos — 1 min','Mostrar onde ficam avisos, oportunidades e agenda.'),('7','Administração — 1 min','Mostrar o painel, cadastro de conteúdo e a regra de acesso administrativo.')])
add_callout(doc,'PRÓXIMO PASSO OPERACIONAL','Cadastrar os primeiros conteúdos, membros e eventos reais antes da demonstração. A estrutura está pronta para receber dados oficiais da TIME.',MAGENTA)

# footer
for sec in doc.sections:
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('TIME Escola de Negócios  |  Guia visual para sócios'); run_style(r,8.5,MUTED)

doc.core_properties.title='TIME Escola de Negócios — Guia Visual para Sócios'
doc.core_properties.author='TIME Escola de Negócios'
doc.save(OUT)
print(OUT)
